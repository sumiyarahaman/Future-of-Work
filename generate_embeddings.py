"""
INST309 Step 1: Sentence Embedding Generation & Cosine Similarity Pipeline

This script converts cleaned job posting descriptions from Excel dataset into 
384-dimensional dense vector embeddings using Hugging Face's sentence-transformers model 


Requirements:
    pip install sentence-transformers pandas openpyxl scikit-learn numpy

Usage:
    python generate_embeddings.py --excel_path inst309_cleaned_roles_dataset_v2.xlsx --top_k 5
"""

import argparse
import sys
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl
from sklearn.metrics.pairwise import cosine_similarity


def load_dataset(excel_path: str) -> pd.DataFrame:
    """Reads and parses the 'Cleaned Roles Dataset' sheet from the Excel file."""
    excel_file = Path(excel_path)
    if not excel_file.exists():
        raise FileNotFoundError(f"Excel file not found at path: {excel_path}")

    wb = openpyxl.load_workbook(excel_file, data_only=True)
    if "Cleaned Roles Dataset" not in wb.sheetnames:
        raise ValueError("Sheet 'Cleaned Roles Dataset' not found in workbook.")

    ws = wb["Cleaned Roles Dataset"]

    roles = []
    for row in range(5, ws.max_row + 1):
        role_id = ws.cell(row, 1).value
        title = ws.cell(row, 2).value
        domain = ws.cell(row, 3).value
        company = ws.cell(row, 4).value
        desc = ws.cell(row, 5).value
        skills = ws.cell(row, 6).value
        aliases = ws.cell(row, 8).value

        if role_id and title and desc:
            combined_text = (
                f"Role Title: {title}. "
                f"Domain: {domain or 'N/A'}. "
                f"Company: {company or 'N/A'}. "
                f"Description: {desc}. "
                f"Key Skills: {skills or 'N/A'}. "
                f"Emerging Aliases: {aliases or 'N/A'}"
            )
            roles.append({
                "id": str(role_id).strip(),
                "title": str(title).strip(),
                "domain": str(domain).strip() if domain else "Uncategorized",
                "company": str(company).strip() if company else "N/A",
                "description": str(desc).strip(),
                "skills": str(skills).strip() if skills else "",
                "aliases": str(aliases).strip() if aliases else "",
                "combined_text": combined_text
            })

    df = pd.DataFrame(roles)
    print(f"Successfully loaded {len(df)} job postings from {excel_file.name}.")
    return df


def generate_embeddings(
    df: pd.DataFrame,
    model_name: str = "sentence-transformers/all-MiniLM-L6-v2"
) -> np.ndarray:
    """Loads SentenceTransformer model and generates normalized L2 vector embeddings."""
    try:
        from sentence_transformers import SentenceTransformer
    except ImportError:
        print("ERROR: 'sentence-transformers' library is not installed.")
        print("Please install it via: pip install sentence-transformers")
        sys.exit(1)

    print(f"\nLoading model '{model_name}'...")
    model = SentenceTransformer(model_name)

    print(f"Generating embeddings for {len(df)} postings...")
    # normalize_embeddings=True ensures vectors have L2 norm = 1.0 (enables fast dot-product cosine similarity)
    embeddings = model.encode(
        df["combined_text"].tolist(),
        show_progress_bar=True,
        normalize_embeddings=True
    )

    print(f"Embeddings matrix shape: {embeddings.shape} (Count, Vector Dimension)")
    return embeddings


def compute_and_display_similarity(
    df: pd.DataFrame,
    embeddings: np.ndarray,
    target_role_id: str = "ROLE-001",
    top_k: int = 5
) -> pd.DataFrame:
    """Computes pairwise cosine similarity matrix and prints top-K matches for a target role."""
    sim_matrix = cosine_similarity(embeddings)

    if target_role_id not in df["id"].values:
        print(f"Warning: Target role ID '{target_role_id}' not found. Defaulting to first entry.")
        target_idx = 0
    else:
        target_idx = df[df["id"] == target_role_id].index[0]

    target_role = df.iloc[target_idx]
    target_vector = embeddings[target_idx]

    print("\n" + "=" * 70)
    print(f"LIVE VECTOR EMBEDDING DEMO: {target_role['id']} - {target_role['title']}")
    print("=" * 70)
    print(f"Organization: {target_role['company']}")
    print(f"Domain:       {target_role['domain']}")
    print(f"Vector Dimensions: {target_vector.shape[0]}")
    print(f"Vector L2 Norm:   {np.linalg.norm(target_vector):.4f}")
    print("\nFirst 12 Floating-Point Dimensions:")
    print(np.round(target_vector[:12], 5))

    print("\n" + "=" * 70)
    print(f"TOP {top_k} SEMANTIC MATCHES FOR: {target_role['title']}")
    print("=" * 70)

    scores = list(enumerate(sim_matrix[target_idx]))
    sorted_scores = sorted(scores, key=lambda x: x[1], reverse=True)

    results = []
    for rank, (idx, score) in enumerate(sorted_scores[: top_k + 1], 1):
        matched = df.iloc[idx]
        is_target = " (Target)" if idx == target_idx else ""
        print(
            f"{rank}. [{matched['id']}] {matched['title'][:35]:<35} | "
            f"Domain: {matched['domain'][:22]:<22} | "
            f"Similarity: {score:.4f}{is_target}"
        )
        results.append({
            "rank": rank,
            "target_id": target_role["id"],
            "matched_id": matched["id"],
            "matched_title": matched["title"],
            "matched_domain": matched["domain"],
            "cosine_similarity": float(score)
        })

    return pd.DataFrame(results)


def save_artifacts(
    df: pd.DataFrame,
    embeddings: np.ndarray,
    output_dir: str = "output"
) -> None:
    """Exports generated embeddings (.npy) and metadata (.csv) for Step 2 clustering."""
    out_path = Path(output_dir)
    out_path.mkdir(parents=True, exist_ok=True)

    # Save numpy embedding matrix
    npy_file = out_path / "inst309_role_embeddings.npy"
    np.save(npy_file, embeddings)
    print(f"\n[Artifact Saved] Vectors exported to: {npy_file}")

    # Save metadata mapping file
    meta_file = out_path / "inst309_roles_metadata.csv"
    df.to_csv(meta_file, index=False)
    print(f"[Artifact Saved] Metadata exported to: {meta_file}")


def main():
    parser = argparse.ArgumentParser(
        description="INST309 Step 1: Sentence Embedding Generation & Cosine Similarity"
    )
    parser.add_argument(
        "--excel_path",
        type=str,
        default="cleaned_roles_dataset.xlsx",
        help="Path to cleaned Excel workbook"
    )
    parser.add_argument(
        "--model_name",
        type=str,
        default="sentence-transformers/all-MiniLM-L6-v2",
        help="Hugging Face model checkpoint"
    )
    parser.add_argument(
        "--target_id",
        type=str,
        default="ROLE-001",
        help="Role ID to demonstrate similarity search"
    )
    parser.add_argument(
        "--top_k",
        type=int,
        default=5,
        help="Number of similar roles to retrieve"
    )
    parser.add_argument(
        "--output_dir",
        type=str,
        default="output",
        help="Directory to save output numpy vectors and CSV metadata"
    )

    args = parser.parse_args()

    # Pipeline Execution
    df = load_dataset(args.excel_path)
    embeddings = generate_embeddings(df, model_name=args.model_name)
    compute_and_display_similarity(
        df, embeddings, target_role_id=args.target_id, top_k=args.top_k
    )
    save_artifacts(df, embeddings, output_dir=args.output_dir)
    print("\nPipeline execution complete! Ready for Step 2 (HDBSCAN Clustering).")


if __name__ == "__main__":
    main()
